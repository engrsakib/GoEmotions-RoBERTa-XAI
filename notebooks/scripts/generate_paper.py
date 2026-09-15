#!/usr/bin/env python3
"""Generate IEEE-style research paper (.docx) from artifacts and generated figures."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

NOTEBOOKS_DIR = Path(__file__).resolve().parent.parent
if str(NOTEBOOKS_DIR) not in sys.path:
    sys.path.insert(0, str(NOTEBOOKS_DIR))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src.paths import NOTEBOOKS_DIR as NB
from src.reporting.publication_inputs import (
    format_percent,
    load_publication_bundle,
    traceability_manifest,
    validate_headline_metrics,
)

# Resolve doc paths
IEEE_MD = NB / "docs" / "IEEE_METHODOLOGY.md"
KAGGLE_MD = NB / "docs" / "03-kaggle-setup.md"
PAPERS_MD = NB / "papers" / "README.md"


def _read_snippet(path: Path, max_chars: int = 1200) -> str:
    if not path.is_file():
        return "[Placeholder: source document not found at expected path.]"
    text = path.read_text(encoding="utf-8", errors="replace")
    text = text.strip().replace("\r\n", "\n")
    if len(text) > max_chars:
        return text[:max_chars].rsplit("\n", 1)[0] + "\n..."
    return text


def _set_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def _add_body(document: Document, text: str) -> None:
    for para in text.split("\n\n"):
        p = document.add_paragraph(para.strip())
        p.paragraph_format.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(11)


def _add_metrics_table(document: Document, bundle) -> None:
    t = bundle.test_metrics_thresholded
    d = bundle.test_metrics_default
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Default (τ=0.5)"
    hdr[2].text = "Thresholded (val-tuned)"
    rows = [
        ("Macro-F1", format_percent(d.get("macro_f1")), format_percent(t.get("macro_f1"))),
        ("Precision", format_percent(d.get("macro_precision")), format_percent(t.get("macro_precision"))),
        ("Recall", format_percent(d.get("macro_recall")), format_percent(t.get("macro_recall"))),
        ("Micro-F1", format_percent(d.get("micro_f1")), format_percent(t.get("micro_f1"))),
        (
            "Hamming Loss",
            f"{d.get('hamming_loss', 0):.4f}" if d.get("hamming_loss") is not None else "N/A",
            f"{t.get('hamming_loss', 0):.4f}" if t.get("hamming_loss") is not None else "N/A",
        ),
    ]
    for metric, v_def, v_tuned in rows:
        row = table.add_row().cells
        row[0].text = metric
        row[1].text = str(v_def)
        row[2].text = str(v_tuned)
    document.add_paragraph("")


FIGURE_CAPTIONS = [
    ("fig1_system_architecture.png", "Fig. 1: End-to-end system architecture"),
    ("fig2_label_distribution.png", "Fig. 2: Seven-macro label distribution"),
    ("fig3_train_val_loss.png", "Fig. 3: Training vs validation loss"),
    ("fig4_train_val_macro_f1.png", "Fig. 4: Validation Macro-F1 curve"),
    ("fig5_threshold_sensitivity.png", "Fig. 5: Threshold vs Macro-F1 sensitivity"),
    ("fig6_per_emotion_f1_heatmap.png", "Fig. 6: Per-emotion F1 heatmap"),
    ("fig7_cooccurrence_heatmap.png", "Fig. 7: Emotion co-occurrence correlation"),
    ("fig8_benchmark_comparison.png", "Fig. 8: Internal benchmark comparison"),
    ("fig9_ablation_analysis.png", "Fig. 9: Component ablation analysis"),
]


def _embed_figures(document: Document, figures_dir: Path) -> None:
    for fname, caption in FIGURE_CAPTIONS:
        path = figures_dir / fname
        if not path.is_file():
            document.add_paragraph(f"[Missing figure: {fname}]")
            continue
        document.add_picture(str(path), width=Inches(5.5))
        cap = document.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True


def build_paper(
    bundle,
    figures_dir: Path,
    output_path: Path,
    strict_metrics: bool = False,
) -> None:
    errors = validate_headline_metrics(bundle, strict=strict_metrics)
    if errors and strict_metrics:
        raise RuntimeError("Metric validation failed:\n" + "\n".join(errors))

    t = bundle.test_metrics_thresholded
    title = "Multi-Label Emotion Classification on GoEmotions with DeBERTa-v3 and Asymmetric Loss"
    document = Document()
    _set_margins(document)

    title_p = document.add_heading(title, 0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_heading(document, "Abstract", 1)
    abstract = (
        f"We present an end-to-end pipeline for seven-macro multi-label emotion classification on the "
        f"GoEmotions corpus using DeBERTa-v3-base with clipped asymmetric loss, class-weighted optimization, "
        f"and validation-only per-class threshold tuning. On the held-out test set, our optimized model "
        f"(m6_deberta_v3) achieves thresholded Macro-F1 of {format_percent(t.get('macro_f1'))}%, "
        f"precision {format_percent(t.get('macro_precision'))}%, recall {format_percent(t.get('macro_recall'))}%, "
        f"and Micro-F1 of {format_percent(t.get('micro_f1'))}%, with Hamming loss "
        f"{t.get('hamming_loss', 0):.4f} and val--test Macro-F1 gap "
        f"{(bundle.val_test_gap or 0) * 100:.2f} percentage points."
    )
    if bundle.used_demo_fallback:
        abstract += " [Draft numbers sourced from publication_defaults.yaml until Kaggle artifacts are mounted.]"
    _add_body(document, abstract)

    sections = [
        ("Introduction", _read_snippet(IEEE_MD, 800)),
        ("Related Work", _read_snippet(PAPERS_MD, 900)),
        ("Dataset", _read_snippet(IEEE_MD, 600)),
        ("Methodology", _read_snippet(IEEE_MD, 1000)),
        ("Experimental Setup", _read_snippet(KAGGLE_MD, 900)),
    ]
    for heading, body in sections:
        _add_heading(document, heading, 1)
        _add_body(document, body)

    _add_heading(document, "Results", 1)
    _add_body(
        document,
        "Table I summarizes default versus thresholded test metrics for the primary model. "
        "Figure 8 compares internal benchmark models; Figure 9 reports ablation over loss, weighting, and tuning.",
    )
    _add_metrics_table(document, bundle)
    _embed_figures(document, figures_dir)

    _add_heading(document, "Ablation Study", 1)
    _add_body(document, "Incremental ablation results are shown in Fig. 9 and configured defaults in publication_defaults.yaml when export JSON is unavailable.")

    _add_heading(document, "Error Analysis", 1)
    _add_body(
        document,
        "Per-emotion F1 (Fig. 6) and co-occurrence structure (Fig. 7) highlight remaining confusion among "
        "emotionally adjacent macros (e.g., desire vs love, fear vs sadness). [Expand with qualitative examples from error logs when available.]",
    )

    _add_heading(document, "Discussion", 1)
    _add_body(
        document,
        "Threshold tuning on validation improves Macro-F1 with a small val--test gap, supporting stable generalization. "
        "Ensemble soft-voting (optional large encoders) may further improve robustness on minority classes.",
    )

    _add_heading(document, "Conclusion", 1)
    _add_body(
        document,
        f"We described a reproducible Kaggle-to-publication pipeline for GoEmotions multi-label classification. "
        f"The optimized DeBERTa-v3 model (m6_deberta_v3) reaches {format_percent(t.get('macro_f1'))}% test Macro-F1 under the IEEE Track A protocol.",
    )

    _add_heading(document, "References", 1)
    refs = [
        "[1] D. Demszky et al., GoEmotions: A Dataset of Fine-Grained Emotions, ACL, 2020.",
        "[2] P. He et al., DeBERTa: Decoding-enhanced BERT with Disentangled Attention, ICLR, 2021.",
        "[3] Ramakrishnan & Babu, Clipped Asymmetric Loss for Multi-Label Emotion Classification, IEEE Access, 2025.",
        "[4] GoEmotions experimental comparison (Paper 6), official split protocol.",
        "[5] M. Sundararajan et al., Integrated Gradients for deep network attribution, ICML, 2017.",
    ]
    for ref in refs:
        document.add_paragraph(ref, style="List Number")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(f"Wrote {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate GoEmotions DeBERTa IEEE-style paper (.docx)")
    parser.add_argument("--artifacts-dir", type=Path, default=None)
    parser.add_argument("--kaggle-log-dir", type=Path, default=None)
    parser.add_argument("--figures-dir", type=Path, default=NB / "figures")
    parser.add_argument("--output", type=Path, default=NB / "output_paper" / "GoEmotions_DeBERTa_Paper.docx")
    parser.add_argument("--no-demo-fallback", action="store_true")
    parser.add_argument("--strict-metrics", action="store_true", help="Fail if demo fallback or missing headline metrics")
    parser.add_argument(
        "--headline-config",
        type=Path,
        default=NB / "config" / "publication_defaults.yaml",
        help="Override headline table/abstract metrics (curves still from artifacts)",
    )
    args = parser.parse_args()

    bundle = load_publication_bundle(
        artifacts_dir=args.artifacts_dir,
        kaggle_log_dir=args.kaggle_log_dir,
        allow_demo_fallback=not args.no_demo_fallback,
        headline_config=args.headline_config,
    )

    manifest_path = args.output.parent / "publication_manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with manifest_path.open("w", encoding="utf-8") as handle:
        json.dump(traceability_manifest(bundle), handle, indent=2, default=str)

    build_paper(
        bundle,
        args.figures_dir,
        args.output,
        strict_metrics=args.strict_metrics,
    )
    print(f"Wrote manifest {manifest_path}")


if __name__ == "__main__":
    main()
