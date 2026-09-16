from pathlib import Path
from textwrap import dedent

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PPTX_PATH = ROOT / "workflow_linear_diagram.pptx"
SVG_PATH = ROOT / "workflow_linear_diagram.svg"
MERMAID_PATH = ROOT / "workflow_linear_diagram.mmd"

SOFT_BLUE = "E6F2FF"
DARK_NAVY = "004080"
WARM_ACCENT = "FFEB3B"
TEXT_DARK = "1F2937"
BG = "FFFFFF"

STAGES = [
    ("Stage 0", "Bootstrap", "Environment setup", False),
    ("Stage 1", "Data Pipeline", "Load, map, clean, split", True),
    ("Stage 2", "EDA", "Distribution checks", False),
    ("Stage 3", "Baselines", "TF-IDF LogReg / SVM", False),
    ("Stage 4", "Train", "Transformer fine-tuning", True),
    ("Stage 5", "Evaluate", "Metrics and reports", False),
    ("Stage 6", "XAI", "Integrated Gradients", True),
    ("Stage 7", "Export", "Saved model package", False),
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def add_textbox(slide, left, top, width, height, text, *, font_size, bold=False, color=TEXT_DARK, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_stage_box(slide, left, top, width, height, stage_no, title, detail, highlight):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(SOFT_BLUE)
    shape.line.color.rgb = rgb(DARK_NAVY)
    shape.line.width = Pt(1.8)

    if highlight:
        badge = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            left + width * 0.29,
            top - Inches(0.22),
            width * 0.42,
            Inches(0.22),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = rgb(WARM_ACCENT)
        badge.line.color.rgb = rgb(DARK_NAVY)
        badge.line.width = Pt(1.0)
        tf = badge.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = "Key Innovation"
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = rgb(DARK_NAVY)

    add_textbox(slide, left + Inches(0.06), top + Inches(0.08), width - Inches(0.12), Inches(0.26), stage_no, font_size=11, bold=True, color=DARK_NAVY)
    add_textbox(slide, left + Inches(0.08), top + Inches(0.36), width - Inches(0.16), Inches(0.34), title, font_size=15, bold=True)
    add_textbox(slide, left + Inches(0.10), top + Inches(0.83), width - Inches(0.20), Inches(0.46), detail, font_size=10.5)
    return shape


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(BG)

    add_textbox(slide, Inches(0.5), Inches(0.25), Inches(15.0), Inches(0.45), "GoEmotions-RoBERTa-XAI Linear Workflow", font_size=23, bold=True, color=TEXT_DARK)
    add_textbox(
        slide,
        Inches(1.2),
        Inches(0.72),
        Inches(13.6),
        Inches(0.35),
        "Simplified straight-line version of the full training and deployment workflow",
        font_size=12,
        color=DARK_NAVY,
    )

    left0 = Inches(0.28)
    top = Inches(2.55)
    width = Inches(1.72)
    height = Inches(1.58)
    gap = Inches(0.16)

    centers = []
    for idx, (stage_no, title, detail, highlight) in enumerate(STAGES):
        left = left0 + idx * (width + gap)
        add_stage_box(slide, left, top, width, height, stage_no, title, detail, highlight)
        centers.append((left + width / 2, top + height / 2))

    for idx in range(len(centers) - 1):
        x1 = left0 + idx * (width + gap) + width
        x2 = left0 + (idx + 1) * (width + gap)
        y = top + height / 2
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y, x2, y)
        connector.line.color.rgb = rgb(DARK_NAVY)
        connector.line.width = Pt(2.2)
        try:
            connector.line.end_arrowhead = True
        except Exception:
            pass

    legend = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(5.05), Inches(4.8), Inches(0.46))
    legend.fill.solid()
    legend.fill.fore_color.rgb = rgb(WARM_ACCENT)
    legend.line.color.rgb = rgb(DARK_NAVY)
    legend.line.width = Pt(1.1)
    tf = legend.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Highlighted blocks mark core method contributions"
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = rgb(DARK_NAVY)

    footer = slide.shapes.add_textbox(Inches(0.5), Inches(8.15), Inches(15.0), Inches(0.3))
    footer_tf = footer.text_frame
    footer_tf.clear()
    footer_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = footer_tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "Editable PowerPoint vector layout generated with python-pptx"
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = rgb(DARK_NAVY)

    prs.save(PPTX_PATH)


def build_svg():
    width = 3000
    height = 720
    box_w = 300
    box_h = 158
    gap = 55
    start_x = 40
    y = 250

    stage_svg = []
    arrow_svg = []
    badge_svg = []

    for idx, (stage_no, title, detail, highlight) in enumerate(STAGES):
        x = start_x + idx * (box_w + gap)
        stage_svg.append(
            f'''
            <rect x="{x}" y="{y}" rx="24" ry="24" width="{box_w}" height="{box_h}"
                  fill="#{SOFT_BLUE}" stroke="#{DARK_NAVY}" stroke-width="4"/>
            <text x="{x + box_w/2}" y="{y + 34}" text-anchor="middle"
                  font-family="Arial, Helvetica, sans-serif" font-size="22" font-weight="700" fill="#{DARK_NAVY}">{stage_no}</text>
            <text x="{x + box_w/2}" y="{y + 76}" text-anchor="middle"
                  font-family="Arial, Helvetica, sans-serif" font-size="25" font-weight="700" fill="#{TEXT_DARK}">{title}</text>
            '''
        )

        detail_lines = detail.split("\n")
        detail_y = y + 116
        for line in detail_lines:
            stage_svg.append(
                f'<text x="{x + box_w/2}" y="{detail_y}" text-anchor="middle" '
                f'font-family="Arial, Helvetica, sans-serif" font-size="18" fill="#{TEXT_DARK}">{line}</text>'
            )
            detail_y += 24

        if highlight:
            badge_svg.append(
                f'''
                <rect x="{x + 76}" y="{y - 34}" rx="14" ry="14" width="148" height="30"
                      fill="#{WARM_ACCENT}" stroke="#{DARK_NAVY}" stroke-width="2"/>
                <text x="{x + box_w/2}" y="{y - 13}" text-anchor="middle"
                      font-family="Arial, Helvetica, sans-serif" font-size="15" font-weight="700" fill="#{DARK_NAVY}">Key Module</text>
                '''
            )

        if idx < len(STAGES) - 1:
            x1 = x + box_w
            x2 = x + box_w + gap - 12
            arrow_svg.append(
                f'<line x1="{x1}" y1="{y + box_h/2}" x2="{x2}" y2="{y + box_h/2}" stroke="#{DARK_NAVY}" stroke-width="4" marker-end="url(#arrowhead)"/>'
            )

    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
  <defs>
    <marker id="arrowhead" markerWidth="14" markerHeight="10" refX="10" refY="5" orient="auto">
      <polygon points="0 0, 10 5, 0 10" fill="#{DARK_NAVY}" />
    </marker>
  </defs>
  <rect width="100%" height="100%" fill="#{BG}"/>
  <text x="{width/2}" y="70" text-anchor="middle" font-family="Arial, Helvetica, sans-serif"
        font-size="40" font-weight="700" fill="#{TEXT_DARK}">GoEmotions-RoBERTa-XAI Linear Workflow</text>
  <text x="{width/2}" y="108" text-anchor="middle" font-family="Arial, Helvetica, sans-serif"
        font-size="22" fill="#{DARK_NAVY}">Simplified single-line layout based on the original stage sequence</text>
  {''.join(badge_svg)}
  {''.join(stage_svg)}
  {''.join(arrow_svg)}
  <rect x="1105" y="540" rx="16" ry="16" width="790" height="46" fill="#{WARM_ACCENT}" stroke="#{DARK_NAVY}" stroke-width="2"/>
  <text x="1500" y="570" text-anchor="middle" font-family="Arial, Helvetica, sans-serif"
        font-size="18" font-weight="700" fill="#{DARK_NAVY}">Highlighted blocks mark core method contributions</text>
</svg>
'''
    SVG_PATH.write_text(svg, encoding="utf-8")


def build_mermaid() -> str:
    mermaid = dedent(
        """
        flowchart LR
            A["Stage 0<br/>Bootstrap<br/>Environment setup"] --> B["Stage 1<br/>Data Pipeline<br/>Load, map, clean, split"]
            B --> C["Stage 2<br/>EDA<br/>Distribution checks"]
            C --> D["Stage 3<br/>Baselines<br/>TF-IDF LogReg / SVM"]
            D --> E["Stage 4<br/>Train<br/>Transformer fine-tuning"]
            E --> F["Stage 5<br/>Evaluate<br/>Metrics and reports"]
            F --> G["Stage 6<br/>XAI<br/>Integrated Gradients"]
            G --> H["Stage 7<br/>Export<br/>Saved model package"]

            classDef process fill:#E6F2FF,stroke:#004080,stroke-width:2px,color:#1F2937;
            classDef innovation fill:#FFEB3B,stroke:#004080,stroke-width:2px,color:#1F2937;

            class A,C,D,F,H process;
            class B,E,G innovation;
        """
    ).strip()
    MERMAID_PATH.write_text(mermaid + "\n", encoding="utf-8")
    return mermaid


def main():
    build_pptx()
    build_svg()
    mermaid = build_mermaid()
    print(f"Wrote {PPTX_PATH}")
    print(f"Wrote {SVG_PATH}")
    print(f"Wrote {MERMAID_PATH}")
    print("\nMermaid source:\n")
    print(mermaid)


if __name__ == "__main__":
    main()
